import io

import pandas as pd
from docx import Document as DocxDocument
from pypdf import PdfReader

IMAGE_EXTENSIONS = {"jpg", "jpeg", "png", "tiff", "tif", "bmp", "webp"}
SUPPORTED_EXTENSIONS = {"pdf", "docx", "xlsx", "xls", "csv", "txt", "md"} | IMAGE_EXTENSIONS

# Minimum characters expected per page in a text-based PDF.
# Fewer than this triggers OCR fallback (scanned PDF).
_MIN_CHARS_PER_PAGE = 40

# Lazy singleton — EasyOCR is expensive to initialise (~2-5 s, loads ~300 MB of models).
_ocr_reader = None
_OCR_MODEL_DIR = "/ocr_models"


class ParseError(Exception):
    pass


def extension_of(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def extract_text(filename: str, content: bytes) -> str:
    ext = extension_of(filename)
    try:
        if ext == "pdf":
            return _extract_pdf(content)
        if ext == "docx":
            return _extract_docx(content)
        if ext in ("xlsx", "xls"):
            return _extract_excel(content)
        if ext == "csv":
            return _extract_csv(content)
        if ext in ("txt", "md"):
            return content.decode("utf-8", errors="ignore")
        if ext in IMAGE_EXTENSIONS:
            return _extract_image(content)
    except ParseError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"خطا در خواندن فایل: {exc}") from exc
    raise ParseError(f"فرمت فایل .{ext} پشتیبانی نمی‌شود")


# ---------- text-based formats ----------

def _extract_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n\n".join(pages).strip()

    # Scanned PDF: if extracted text is too sparse, fall back to OCR
    page_count = max(len(reader.pages), 1)
    if len(text) < _MIN_CHARS_PER_PAGE * page_count:
        ocr_text = _ocr_pdf(content)
        return ocr_text if ocr_text else text
    return text


def _extract_docx(content: bytes) -> str:
    doc = DocxDocument(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _extract_excel(content: bytes) -> str:
    sheets = pd.read_excel(io.BytesIO(content), sheet_name=None, dtype=str)
    parts = []
    for name, df in sheets.items():
        df = df.fillna("")
        parts.append(f"# شیت: {name}\n{df.to_csv(index=False)}")
    return "\n\n".join(parts).strip()


def _extract_csv(content: bytes) -> str:
    df = pd.read_csv(io.BytesIO(content), dtype=str).fillna("")
    return df.to_csv(index=False).strip()


# ---------- OCR ----------

def _get_ocr_reader():
    """Lazy singleton for EasyOCR — initialised on first OCR call."""
    global _ocr_reader
    if _ocr_reader is None:
        try:
            import easyocr  # noqa: PLC0415
            _ocr_reader = easyocr.Reader(
                ["fa", "en", "ar"],
                gpu=False,
                verbose=False,
                model_storage_directory=_OCR_MODEL_DIR,
            )
        except ImportError as exc:
            raise ParseError("کتابخانه EasyOCR نصب نشده است") from exc
    return _ocr_reader


def _run_ocr(img_bytes: bytes) -> str:
    """Run OCR on raw image bytes; returns extracted text."""
    import numpy as np  # noqa: PLC0415
    from PIL import Image as PILImage  # noqa: PLC0415

    reader = _get_ocr_reader()
    img = PILImage.open(io.BytesIO(img_bytes)).convert("RGB")
    results = reader.readtext(np.array(img), detail=1)
    lines = [text for _, text, conf in results if conf > 0.3]
    return "\n".join(lines)


def _extract_image(content: bytes) -> str:
    text = _run_ocr(content)
    if not text.strip():
        raise ParseError("متنی در تصویر شناسایی نشد")
    return text


def _ocr_pdf(content: bytes) -> str:
    """Convert each page of a scanned PDF to an image and OCR it."""
    try:
        import fitz  # PyMuPDF  # noqa: PLC0415
    except ImportError as exc:
        raise ParseError("کتابخانه PyMuPDF نصب نشده است") from exc

    doc = fitz.open(stream=content, filetype="pdf")
    parts = []
    for page in doc:
        pix = page.get_pixmap(dpi=150)
        page_text = _run_ocr(pix.tobytes("png"))
        if page_text.strip():
            parts.append(page_text)
    return "\n\n".join(parts).strip()
